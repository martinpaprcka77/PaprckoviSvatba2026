"""Rebuild rozpocet-svatba-2026.xlsx from data/tasks.md."""
import re, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent.parent
TASKS_MD = ROOT / 'data' / 'tasks.md'
OUTPUT = ROOT / 'data' / 'rozpocet-svatba-2026.xlsx'

BLUE = Font(color='0000FF')
BLACK = Font(color='000000')
BOLD = Font(bold=True)
BOLD_BLACK = Font(bold=True, color='000000')
HEADER_FILL = PatternFill('solid', fgColor='722F37')
HEADER_FONT = Font(bold=True, color='FFFFFF')
RED_FILL = PatternFill('solid', fgColor='FFC7CE')
GREEN_FILL = PatternFill('solid', fgColor='C6EFCE')
YELLOW_FILL = PatternFill('solid', fgColor='FFEB9C')
GRAY_FILL = PatternFill('solid', fgColor='D9D9D9')
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'))

def parse_tasks():
    md = TASKS_MD.read_text(encoding='utf-8')
    tasks = []
    in_table = False
    for line in md.split('\n'):
        t = line.strip()
        if t.startswith('| Termín') and 'Úkol' in t: in_table = True; continue
        if t.startswith('|--'): continue
        if not in_table or not t.startswith('|'): continue
        if t.startswith('## ') or t.startswith('# '): in_table = False; continue
        cells = [c.strip() for c in t.split('|') if c.strip() != '']
        if len(cells) < 9: continue
        termin, ukol, kdo, kat, plan_str, skut_str, stav, timestamp, poznamka = cells
        plan = int(plan_str.replace(' ', '')) if plan_str.replace(' ', '').replace('_','').isdigit() else 0
        skut = int(skut_str.replace(' ', '')) if skut_str not in ('_', '') and skut_str.replace(' ','').isdigit() else None
        done = '[x]' in stav
        tasks.append({'termin': termin, 'ukol': ukol, 'kdo': kdo, 'kat': kat,
                       'plan': plan, 'skut': skut, 'done': done,
                       'timestamp': timestamp if timestamp != '_' else '',
                       'poznamka': poznamka if poznamka != '_' else ''})
    return tasks

def build():
    tasks = parse_tasks()
    wb = Workbook()

    # Sheet 1: Úkoly
    ws = wb.active
    ws.title = 'Úkoly'
    headers = ['ID', 'Termín', 'Úkol', 'Kdo', 'Kategorie', 'Plán (Kč)', 'Skutečnost (Kč)', 'Stav', 'Timestamp', 'Poznámka']
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        cell.border = THIN_BORDER
    for r, t in enumerate(tasks, 2):
        vals = [r-1, t['termin'], t['ukol'], t['kdo'], t['kat'],
                t['plan'], t['skut'] if t['skut'] else '',
                '✓' if t['done'] else ' ', t['timestamp'], t['poznamka']]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = THIN_BORDER
            if c == 6:  # Plan - blue input
                cell.font = BLUE; cell.number_format = '#,##0'
            elif c == 7:  # Actual - blue input
                cell.font = BLUE; cell.number_format = '#,##0'
            else:
                cell.font = Font()
        if t['done']:
            for c in range(1, 11):
                ws.cell(row=r, column=c).fill = GRAY_FILL
                ws.cell(row=r, column=c).font = Font(color='808080')
    ws.column_dimensions['A'].width = 5; ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 50; ws.column_dimensions['D'].width = 22
    ws.column_dimensions['E'].width = 14; ws.column_dimensions['F'].width = 14
    ws.column_dimensions['G'].width = 16; ws.column_dimensions['H'].width = 8
    ws.column_dimensions['I'].width = 20; ws.column_dimensions['J'].width = 30
    ws.auto_filter.ref = f'A1:J{len(tasks)+1}'
    ws.freeze_panes = 'A2'

    # Sheet 2: Rozpočet
    ws2 = wb.create_sheet('Rozpočet')
    cat_labels = {'mandatory': 'Povinné', 'important': 'Důležité', 'optional': 'Volitelné'}
    ws2.cell(row=1, column=1, value='Kategorie').font = HEADER_FONT
    ws2.cell(row=1, column=1).fill = HEADER_FILL; ws2.cell(row=1, column=1).border = THIN_BORDER
    ws2.cell(row=1, column=2, value='Počet úkolů').font = HEADER_FONT
    ws2.cell(row=1, column=2).fill = HEADER_FILL; ws2.cell(row=1, column=2).border = THIN_BORDER
    ws2.cell(row=1, column=3, value='Plán (Kč)').font = HEADER_FONT
    ws2.cell(row=1, column=3).fill = HEADER_FILL; ws2.cell(row=1, column=3).border = THIN_BORDER
    ws2.cell(row=1, column=4, value='Skutečnost (Kč)').font = HEADER_FONT
    ws2.cell(row=1, column=4).fill = HEADER_FILL; ws2.cell(row=1, column=4).border = THIN_BORDER

    cats = ['mandatory', 'important', 'optional']
    for i, cat in enumerate(cats):
        r = i + 2
        ws2.cell(row=r, column=1, value=cat_labels[cat]).border = THIN_BORDER
        ws2.cell(row=r, column=2, value=f'=COUNTIF(Úkoly!E2:E{len(tasks)+1},"{cat}")').border = THIN_BORDER
        ws2.cell(row=r, column=3, value=f'=SUMIF(Úkoly!E2:E{len(tasks)+1},"{cat}",Úkoly!F2:F{len(tasks)+1})').border = THIN_BORDER
        ws2.cell(row=r, column=3).number_format = '#,##0'
        ws2.cell(row=r, column=4, value=f'=SUMIF(Úkoly!E2:E{len(tasks)+1},"{cat}",Úkoly!G2:G{len(tasks)+1})').border = THIN_BORDER
        ws2.cell(row=r, column=4).number_format = '#,##0'

    r_total = 5
    ws2.cell(row=r_total, column=1, value='CELKEM').font = BOLD; ws2.cell(row=r_total, column=1).border = THIN_BORDER
    ws2.cell(row=r_total, column=2, value=f'=SUM(B2:B4)').font = BOLD; ws2.cell(row=r_total, column=2).border = THIN_BORDER
    ws2.cell(row=r_total, column=3, value=f'=SUM(C2:C4)').font = BOLD; ws2.cell(row=r_total, column=3).number_format = '#,##0'
    ws2.cell(row=r_total, column=3).border = THIN_BORDER
    ws2.cell(row=r_total, column=4, value=f'=SUM(D2:D4)').font = BOLD; ws2.cell(row=r_total, column=4).number_format = '#,##0'
    ws2.cell(row=r_total, column=4).border = THIN_BORDER

    ws2.conditional_formatting.add(f'C{r_total}', CellIsRule(operator='greaterThan', formula=['100000'], fill=RED_FILL, font=BOLD))
    ws2.conditional_formatting.add(f'C{r_total}', CellIsRule(operator='lessThanOrEqual', formula=['100000'], fill=GREEN_FILL, font=BOLD))

    ws2.column_dimensions['A'].width = 18; ws2.column_dimensions['B'].width = 16
    ws2.column_dimensions['C'].width = 16; ws2.column_dimensions['D'].width = 18

    # Sheet 3: Na osobu
    ws3 = wb.create_sheet('Na osobu')
    people = ['Mamka', 'Taťka', 'Žanetka', 'Kikinka', 'Děti']
    for c, h in enumerate(['Osoba', 'Počet úkolů', 'Plán (Kč)', 'Hotovo'], 1):
        cell = ws3.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT; cell.fill = HEADER_FILL; cell.border = THIN_BORDER

    for i, person in enumerate(people):
        r = i + 2
        ws3.cell(row=r, column=1, value=person).border = THIN_BORDER
        ws3.cell(row=r, column=1).font = BOLD
        ws3.cell(row=r, column=2, value=f'=COUNTIF(Úkoly!D2:D{len(tasks)+1},"*{person}*")').border = THIN_BORDER
        ws3.cell(row=r, column=3, value=f'=SUMIF(Úkoly!D2:D{len(tasks)+1},"*{person}*",Úkoly!F2:F{len(tasks)+1})').border = THIN_BORDER
        ws3.cell(row=r, column=3).number_format = '#,##0'
        ws3.cell(row=r, column=4, value=f'=SUMPRODUCT((ISNUMBER(SEARCH("{person}",Úkoly!D2:D{len(tasks)+1})))*(Úkoly!H2:H{len(tasks)+1}="✓"))').border = THIN_BORDER

    ws3.column_dimensions['A'].width = 18; ws3.column_dimensions['B'].width = 16
    ws3.column_dimensions['C'].width = 14; ws3.column_dimensions['D'].width = 12

    wb.save(OUTPUT)
    print(f'OK: {OUTPUT.name} — {len(tasks)} tasks')

if __name__ == '__main__':
    build()
